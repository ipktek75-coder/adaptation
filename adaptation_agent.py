# -*- coding: utf-8 -*-
"""
adaptation_agent.py — АГЕНТ ФОРМИРОВАНИЯ ИНДИВИДУАЛЬНОГО ПЛАНА АДАПТАЦИИ (до 3 месяцев).

Что делает (по ТЗ и РИ 7.2-01-2024 «Адаптация персонала»):
  1. Сканирует папку «Сотрудники» и находит карточки с профилем должности в ЛЮБОМ
     поддерживаемом формате: PDF, DOCX, DOC, RTF, ODT, XLSX, XLS, TXT, MD,
     а также сканы/фото (PNG, JPG, JPEG, BMP, TIF, TIFF, WEBP);
  2. Извлекает текст по формату: текстовый слой PDF (PyMuPDF), иначе OCR (EasyOCR
     ru+en — Tesseract в этом окружении не установлен); DOCX/ODT/XLSX/TXT/MD
     читаются напрямую (python-docx + стандартная библиотека), DOC/RTF и XLS —
     через установленные Word/Excel (pywin32), изображения — через OCR;
  3. Через локальную LLM (Ollama, по умолчанию qwen3.5-64k:latest) разбирает профиль:
     Ф.И.О., подразделение, должность, руководитель, наставник, менеджер программ
     обучения, оценки по блокам «4.1 Теоретические знания» и «4.2 Практические
     навыки» — и находит позиции НИЖЕ нормативных (оценка < 4 из 5 либо пусто);
  4. Если данных о знаниях/умениях мало — ищет в интернете профстандарт/ЕТКС
     по должности (ddgs + загрузка страницы) и берёт требования оттуда;
  5. Собирает 3 блока плана (общий, индивидуальный, итоговое индивидуальное
     практическое задание), упорядочивает по срокам (окна дат считает код
     от срока адаптации) и заполняет Word-шаблон «2_РИ_План адаптации_шаблон.docx»;
  6. Должности из Перечня (Приложение А РИ) — план НЕ формируется, в отчёт
     пишется JSON с "не требуется": true;
  7. ВСЕ поля плана обязательны, КРОМЕ столбца «Подпись адаптируемого
     работника» (он остаётся пустым — подписывается от руки). Если данные по
     полю получить не удалось — в документе выводится КРАСНЫМ
     «не удалось получить данные». В колонке «Участники» всегда
     «Должность — Фамилия И.О.». Полное Ф.И.О. в шапке — ТОЛЬКО у адаптируемого
     работника, остальные строки шапки — «Фамилия И.О.». Подписи под таблицами —
     в формате «Образец.docx»: строка
     «<Должность>  _______________  Фамилия И.О.  <дата>» и слово «Подпись»
     (10 пт) по центру линии; колонки выравниваются таб-стопами
     7,54 / 9,55 / 11,56 / 14,74 см. Пояснения «(для руководителей, принятых в
     дирекцию)» и «(для остальных работников организации)» (9 пт) печатаются на
     строке со словом «Подпись» слева (на уровне надписи должности), «Подпись»
     при этом не сдвигается; после блока «Генеральный директор» — пустая строка.
     Линия подписи гарантированно есть в каждой строке, в том числе у наставника,
     когда наставник не назначен (ensure_sign_line).
     Постоянные подписанты: менеджер программ
     обучения — Смирнова Светлана Андреевна, генеральный директор —
     Зубарев Ю.П., руководитель УДЦ — Баталова Т.В.;
  8. Шрифт: Times New Roman 11 пт — в шапке (таблица 1) и в таблице мероприятий
     (кроме строки с названиями столбцов); в нижней части под таблицами все
     размеры 12 пт заменены на 11 пт (9 и 10 пт сохраняются).

Запуск:
    python adaptation_agent.py                       # все карточки из «Сотрудники»
    python adaptation_agent.py "путь\\к\\ПД.docx"    # одна карточка (любой формат)
    python adaptation_agent.py "путь\\к\\папке" -r   # папка, включая подпапки
    python adaptation_agent.py --model qwen3.5-64k:latest --no-search
    python adaptation_agent.py --ocr-only             # только извлечь и показать текст

Зависимости: python-docx, PyMuPDF, easyocr, ollama, ddgs, requests, pillow, numpy;
для DOC/RTF и XLS дополнительно нужны Word/Excel + pywin32 (необязательно).
"""
from __future__ import annotations

import argparse
import copy
import io
import json
import os
import re
import sys
import time
from datetime import datetime, timedelta

# --- НАСТРОЙКИ ПУТЕЙ ---
BASE = r"C:\Users\pzbayur\Desktop\Проработка Адаптация"
SRC_FOLDER = os.path.join(BASE, "Сотрудники")
OUT_FOLDER = os.path.join(BASE, "План адаптации")
TEMPLATE = os.path.join(BASE, "2_РИ_План адаптации_шаблон.docx")
EXCLUDE_DOC = os.path.join(
    BASE, "2_РИ_Перечень работников, которым не проводится адаптация.docx")
CACHE_FOLDER = os.path.join(OUT_FOLDER, "_cache")
LOG_PATH = os.path.join(OUT_FOLDER, "adaptation_agent.log")

# --- LLM: локальная модель Qwen3.5-64k через Ollama ---
# ADAPT_MODEL — имя модели; ADAPT_OLLAMA_HOST / OLLAMA_HOST — адрес Ollama
# (например, http://192.168.1.50:11434 — сервер Ollama в локальной сети).
MODEL = os.environ.get("ADAPT_MODEL", "qwen3.5-64k:latest")
OLLAMA_HOST = (os.environ.get("ADAPT_OLLAMA_HOST")
               or os.environ.get("OLLAMA_HOST") or "")
LLM_TEMPERATURE = 0.1
LLM_NUM_CTX = 16384
# Qwen3.5-64k по умолчанию «размышляет» (thinking): на JSON-план это лишние
# десятки секунд и риск обрыва ответа. Отключаем; включается ADAPT_THINK=1.
LLM_THINK = os.environ.get("ADAPT_THINK", "0").strip().lower() in (
    "1", "true", "yes", "on")

# --- Постоянные подписанты плана (подставляются во ВСЕ планы) ---
# Переопределяются env-переменными или полями JSON-профиля.
GENERAL_DIRECTOR = os.environ.get("ADAPT_GENERAL_DIRECTOR", "Зубарев Ю.П.")
UDC_HEAD = os.environ.get("ADAPT_UDC_HEAD", "Баталова Т.В.")
PROGRAM_MANAGER = os.environ.get("ADAPT_PROGRAM_MANAGER",
                                 "Смирнова Светлана Андреевна")

# --- Строки подписей под таблицами (форматирование взято из «Образец.docx»)
# Колонки «линия для подписи | Фамилия И.О. | дата» выравниваются ЯВНЫМИ
# таб-стопами (7.54 / 9.55 / 11.56 / 14.74 см) — во всех строках они стоят на
# одной горизонтали; под линией печатается слово «Подпись» (10 пт).
SIGN_TAB_FRACS = (0.45, 0.57, 0.69, 0.88)
SIGN_LINE = "_______________"          # линия для подписи от руки
SIGN_WORD = "Подпись"                   # подсказка по центру линии
FONT_NAME = "Times New Roman"           # единый шрифт документа
FONT_SIZE_PT = 11                       # основной текст: шапка, таблица, подписи
FONT_COLOR = None                       # ленивый RGBColor(0, 0, 0)
SIGN_WORD_PT = 10                       # размер слова «Подпись»
SIGN_NOTE_PT = 9                        # размер пояснений «(для …)»


def font_color():
    """Явный чёрный цвет шрифта.

    В шаблоне у части строк цвет не задан (наследуется «авто»): из-за этого в
    некоторых темах/просмотрщиках Word линия подписи может не отображаться.
    Поэтому во всех элементах нижней части ставим чёрный цвет явно.
    """
    global FONT_COLOR
    if FONT_COLOR is None:
        from docx.shared import RGBColor
        FONT_COLOR = RGBColor(0x00, 0x00, 0x00)
    return FONT_COLOR
# ширина линии подписи (15 знаков «_» ≈ 0,5 em каждый), см
SIGN_LINE_WIDTH_CM = round(len(SIGN_LINE) * 0.5 * FONT_SIZE_PT / 72.0 * 2.54, 2)
# Пояснения к блокам «Утверждаю» — печатаются на строке со словом «Подпись»
# (слева, на уровне надписи должности), слово «Подпись» при этом не сдвигается.
NOTE_GD = "(для руководителей, принятых в дирекцию)"
NOTE_UDC = "(для остальных работников организации)"

# --- OCR (настройки взяты из проверенного process_checklists.py) ---
OCR_LANGS = ("ru", "en")
OCR_SCALE = 2.0
OCR_CONTRAST = 1.8
OCR_LINE_TOL = 18          # склейка боксов в строку, px
OCR_MIN_CONF = 0.30        # порог доверия к распознанному фрагменту

# --- Поддерживаемые форматы карточек в папке «Сотрудники» ---
# PDF: текстовый слой (PyMuPDF) -> OCR; Office/текст: чтение напрямую;
# картинки (сканы/фото): OCR EasyOCR.
PDF_EXTENSIONS = (".pdf",)
DOCX_EXTENSIONS = (".docx", ".docm", ".dotx")
ODT_EXTENSIONS = (".odt",)                      # zip + content.xml (stdlib)
PLAIN_EXTENSIONS = (".txt", ".md", ".text", ".csv")   # plain text
WORD_EXTENSIONS = (".doc", ".rtf")              # через Word (pywin32)
XLSX_EXTENSIONS = (".xlsx", ".xlsm")            # zip + XML (stdlib)
XLS_EXTENSIONS = (".xls",)                      # через Excel (pywin32)
IMAGE_EXTENSIONS = (".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp")
CARD_EXTENSIONS = (PDF_EXTENSIONS + DOCX_EXTENSIONS + ODT_EXTENSIONS
                   + PLAIN_EXTENSIONS + WORD_EXTENSIONS + XLSX_EXTENSIONS
                   + XLS_EXTENSIONS + IMAGE_EXTENSIONS)
# Расширение -> человекочитаемое название (для логов и сообщений об ошибках)
EXT_NAMES = {
    ".pdf": "PDF", ".docx": "DOCX", ".docm": "DOCM", ".dotx": "DOTX",
    ".doc": "DOC (Word)", ".rtf": "RTF", ".odt": "ODT",
    ".xlsx": "XLSX", ".xlsm": "XLSM", ".xls": "XLS (Excel)",
    ".txt": "TXT", ".md": "MD", ".text": "TXT", ".csv": "CSV",
    ".png": "PNG", ".jpg": "JPG", ".jpeg": "JPEG", ".bmp": "BMP",
    ".tif": "TIF", ".tiff": "TIFF", ".webp": "WEBP",
}

# --- Перечень работников, которым адаптация НЕ проводится (Приложение А РИ).
#     Если рядом есть docx «2_РИ_Перечень...», список читается из него. ---
EXCLUDE_FALLBACK = (
    "Бригадир подсобных рабочих",
    "Подсобный рабочий",
    "Уборщик производственных и служебных помещений",
    "Водитель автомобиля",
    "Тракторист",
    "Транспортировщик",
    "Транспортный рабочий",
    "Распределитель работ",
    "Лифтер",
    "Секретарь",
)

# --- Сроки (как в примере плана: общий блок — первая неделя,
#     индивидуальный — середина, итоговое задание — последние 4 недели) ---
GENERAL_DAYS = 7
FINAL_DAYS = 28

LOG_LINES: list[str] = []


def log(msg: str, tag: str = "АГЕНТ") -> None:
    """Печать в консоль, накопление и немедленная запись строки в лог-файл."""
    line = f"[{tag}] {msg}"
    print(line, flush=True)
    LOG_LINES.append(line)
    try:                                  # инкрементально, чтобы видеть прогресс
        os.makedirs(OUT_FOLDER, exist_ok=True)
        with open(LOG_PATH, "a", encoding="utf-8") as f:
            f.write(line + "\n")
    except Exception:
        pass


def save_log() -> None:
    os.makedirs(OUT_FOLDER, exist_ok=True)
    with open(LOG_PATH, "w", encoding="utf-8") as f:
        f.write("\n".join(LOG_LINES) + "\n")


def norm(text: str) -> str:
    """Нормализация строки для сравнения должностей."""
    t = (text or "").lower().replace("ё", "е")
    t = re.sub(r"[^а-яa-z0-9]+", " ", t)
    return " ".join(t.split())


# ----------------------------------------------------------------------
# ОБЯЗАТЕЛЬНОЕ ЗАПОЛНЕНИЕ ПОЛЕЙ ПЛАНА
#   Все поля плана обязательны, КРОМЕ столбца «Подпись адаптируемого
#   работника» (его оставляем пустым — работник подписывает от руки).
#   Если данные по полю получить не удалось — пишем КРАСНЫМ:
#   «не удалось получить данные».
# ----------------------------------------------------------------------
MISSING_TEXT = "не удалось получить данные"
_MISSING_COLOR = None      # docx.shared.RGBColor подставляется лениво


def missing_color():
    """Красный цвет шрифта для незаполненных полей (ленивая загрузка)."""
    global _MISSING_COLOR
    if _MISSING_COLOR is None:
        from docx.shared import RGBColor
        _MISSING_COLOR = RGBColor(0xFF, 0x00, 0x00)
    return _MISSING_COLOR


def is_missing(value) -> bool:
    """Поле не заполнено: пусто либо шаблонная заглушка («Ф. И. О.» и т. п.)."""
    if value is None:
        return True
    t = str(value).strip()
    if not t:
        return True
    n = norm(t)
    return n in ("ф и о", "не удалось получить данные", "нет данных",
                 "неизвестно", "не определено", "не указано", "н д")


def field_line(value):
    """(текст, признак_красного) для поля: значение либо «не удалось получить данные»."""
    if is_missing(value):
        return (MISSING_TEXT, True)
    return (str(value).strip(), False)


def short_fio(fio: str) -> str:
    """«Фамилия Имя Отчество» -> «Фамилия И.О.» (уже сокращённое — как есть)."""
    parts = [p for p in re.split(r"\s+", str(fio or "").strip()) if p]
    if not parts:
        return ""
    if len(parts) == 1:
        return parts[0]
    inits = []
    for p in parts[1:]:
        up = p.upper()
        if "." in p:                       # уже инициалы: «Ю.С.» -> Ю.С.
            letters = re.findall(r"[А-ЯЁA-Z]", up)
        else:                              # полное слово -> только первая буква
            m = re.search(r"[А-ЯЁA-Z]", up)
            letters = [m.group(0)] if m else []
        if letters:
            inits.append("".join(f"{c}." for c in letters))
    return (parts[0] + " " + "".join(inits)).strip()


def read_exclusion_list() -> list[str]:
    """Список должностей из Приложения А (docx) либо встроенный."""
    if os.path.exists(EXCLUDE_DOC):
        try:
            from docx import Document
            doc = Document(EXCLUDE_DOC)
            items = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            items = [i for i in items if len(i) > 3 and "Перечень" not in i]
            if items:
                log(f"список исключений прочитан из "
                    f"«{os.path.basename(EXCLUDE_DOC)}»: {len(items)} должностей",
                    "ИСКЛЮЧЕНИЯ")
                return items
        except Exception as e:
            log(f"не удалось прочитать docx со списком исключений: "
                f"{type(e).__name__}: {e}", "ИСКЛЮЧЕНИЯ")
    log(f"используется встроенный список исключений: "
        f"{len(EXCLUDE_FALLBACK)} должностей", "ИСКЛЮЧЕНИЯ")
    return list(EXCLUDE_FALLBACK)


def is_excluded(position: str, exclude_list: list[str]) -> str | None:
    """Должность в перечне? Возвращает совпавшую строку списка либо None."""
    p = norm(position)
    if not p:
        return None
    for item in exclude_list:
        n = norm(item)
        if not n:
            continue
        if p == n or p.startswith(n) or n in p:
            return item
    return None


def date_range(start: datetime, end: datetime) -> str:
    return f"{start.strftime('%d.%m')}-{end.strftime('%d.%m')}"


def split_period(start: datetime, end: datetime, parts: int) -> list:
    """Разбить период [start..end] на parts последовательных окон (без дыр)."""
    parts = max(1, int(parts))
    total = max(1, (end - start).days)
    step = max(1, total // parts)
    out = []
    cur = start
    for i in range(parts):
        last = (i == parts - 1)
        nxt = end if last else min(end, cur + timedelta(days=step))
        out.append((cur, nxt))
        cur = nxt
    return out


# ----------------------------------------------------------------------
# 1. СКАНИРОВАНИЕ ПАПКИ СОТРУДНИКОВ
# ----------------------------------------------------------------------
def card_ext(path: str) -> str:
    """Расширение файла в нижнем регистре ('.pdf', '.docx', …)."""
    return os.path.splitext(str(path))[1].lower()


def is_supported_card(path: str) -> bool:
    """Файл — карточка поддерживаемого формата (не служебный и не скрытый)."""
    name = os.path.basename(str(path))
    if not name or name.startswith("~$") or name.startswith("."):
        return False                 # временные файлы Word/Excel и скрытые
    return card_ext(name) in CARD_EXTENSIONS


def _natural_key(text: str) -> list:
    """Ключ естественной сортировки: 1, 2, … 10 (а не 1, 10, 2)."""
    parts = re.split(r"(\d+)", str(text).lower())
    return [int(p) if p.isdigit() else p for p in parts]


def scan_cards(folder: str = SRC_FOLDER, recursive: bool = False) -> list:
    """Найти карточки сотрудников (PDF, DOCX/DOCM/DOTX, DOC, RTF, ODT,
    XLSX/XLSM, XLS, TXT/MD/CSV, сканы-картинки) в папке «Сотрудники» в
    естественном порядке имён.

    recursive=True — искать также во вложенных папках.
    """
    if not os.path.isdir(folder):
        log(f"ОШИБКА: папка не найдена: {folder}", "ПОИСК")
        return []
    found = []
    if recursive:
        for root, dirs, names in os.walk(folder):
            dirs.sort(key=_natural_key)
            for name in sorted(names, key=_natural_key):
                if is_supported_card(name):
                    found.append(os.path.join(root, name))
    else:
        for name in sorted(os.listdir(folder), key=_natural_key):
            if is_supported_card(name):
                found.append(os.path.join(folder, name))
    found.sort(key=lambda p: _natural_key(os.path.relpath(p, folder)))

    # не подхватываем уже сформированные планы из папки результатов
    out_abs = os.path.normcase(os.path.abspath(OUT_FOLDER))
    folder_abs = os.path.normcase(os.path.abspath(folder))
    if out_abs != folder_abs and out_abs.startswith(folder_abs + os.sep):
        kept = [p for p in found
                if not os.path.normcase(os.path.abspath(p)).startswith(
                    out_abs + os.sep)]
        if len(kept) != len(found):
            log(f"пропущено файлов из папки результатов "
                f"«{os.path.basename(OUT_FOLDER)}»: {len(found) - len(kept)}",
                "ПОИСК")
            found = kept

    counts = {}
    for path in found:
        ext = card_ext(path)
        counts[ext] = counts.get(ext, 0) + 1
    by_format = ", ".join(f"{EXT_NAMES.get(e, e)}: {n}"
                          for e, n in sorted(counts.items()))
    scope = "с подпапками" if recursive else "без подпапок"
    log(f"в папке «{os.path.basename(folder)}» ({scope}) найдено документов: "
        f"{len(found)}" + (f" — {by_format}" if by_format else ""), "ПОИСК")
    for i, path in enumerate(found, 1):
        rel = os.path.relpath(path, folder)
        size = os.path.getsize(path)
        log(f"  {i}. {rel} [{EXT_NAMES.get(card_ext(path), '?')}] "
            f"({size / 1024 / 1024:.2f} МБ)", "ПОИСК")
    return found


def wait_enter(enabled: bool = True) -> None:
    """Ожидание нажатия Enter, чтобы окно терминала не закрылось."""
    if not enabled:
        return
    try:
        input("\nНажмите Enter для выхода...")
    except (EOFError, KeyboardInterrupt):
        pass


# ----------------------------------------------------------------------
# 2. ИЗВЛЕЧЕНИЕ ТЕКСТА: слой PDF -> OCR (EasyOCR) с кэшем
# ----------------------------------------------------------------------
def _preprocess(img):
    """Подготовка скана к OCR (как в process_checklists.py)."""
    import numpy as np
    from PIL import Image, ImageEnhance, ImageFilter, ImageOps
    im = ImageOps.autocontrast(img.convert("L"))
    im = im.resize((int(im.width * OCR_SCALE), int(im.height * OCR_SCALE)),
                   Image.LANCZOS)
    im = ImageEnhance.Contrast(im).enhance(OCR_CONTRAST)
    im = im.filter(ImageFilter.SHARPEN)
    return np.asarray(im)


def _boxes_to_lines(items) -> list:
    """Собрать читаемые строки из боксов EasyOCR (сортировка по y, затем x)."""
    rows = []
    for box, text, conf in items:
        if conf is not None and conf < OCR_MIN_CONF:
            continue
        ys = [p[1] for p in box]
        xs = [p[0] for p in box]
        rows.append({"y": min(ys), "x": min(xs), "t": text, "c": conf})
    rows.sort(key=lambda r: (r["y"], r["x"]))
    lines, cur = [], []
    for r in rows:
        if cur and abs(r["y"] - cur[0]["y"]) > OCR_LINE_TOL:
            lines.append(cur)
            cur = []
        cur.append(r)
    if cur:
        lines.append(cur)
    out = []
    for line in lines:
        line.sort(key=lambda r: r["x"])
        out.append("  ".join(r["t"] for r in line))
    return out


_OCR_READER = None


def ocr_reader():
    """Единый экземпляр EasyOCR: модель загружается один раз за весь запуск."""
    global _OCR_READER
    if _OCR_READER is None:
        from easyocr import Reader
        try:
            _OCR_READER = Reader(lang_list=list(OCR_LANGS), gpu=False,
                                 verbose=False)
        except TypeError:                 # другие версии EasyOCR
            _OCR_READER = Reader(langs=list(OCR_LANGS), gpu=False, verbose=False)
    return _OCR_READER


def _ocr_cache_path(path: str) -> str:
    """Путь к JSON-кэшу OCR для карточки любого формата."""
    os.makedirs(CACHE_FOLDER, exist_ok=True)
    return os.path.join(CACHE_FOLDER,
                        os.path.splitext(os.path.basename(path))[0] + "_ocr.json")


def _load_ocr_cache(path: str):
    """Прочитать кэш OCR по карточке (или None, если кэша нет)."""
    cache = _ocr_cache_path(path)
    if not os.path.exists(cache):
        return None
    try:
        with open(cache, encoding="utf-8") as f:
            data = json.load(f)
        log(f"OCR взят из кэша: {os.path.basename(cache)} "
            f"(страниц {len(data.get('pages', {}))})", "OCR")
        return data
    except Exception as e:
        log(f"кэш OCR не прочитан ({type(e).__name__}: {e}) — распознаю заново",
            "OCR")
        return None


def _save_ocr_cache(path: str, data: dict) -> None:
    cache = _ocr_cache_path(path)
    with open(cache, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    log(f"кэш OCR: {os.path.basename(cache)}", "OCR")


def ocr_pdf(pdf_path: str) -> dict:
    """OCR сканированного PDF. Результат кэшируется в JSON."""
    import pymupdf
    from PIL import Image
    data = _load_ocr_cache(pdf_path)
    if data is not None:
        return data
    reader = ocr_reader()
    doc = pymupdf.open(pdf_path)
    pages = {}
    t0 = time.time()
    for pno in range(len(doc)):
        page = doc[pno]
        pix = page.get_pixmap(matrix=pymupdf.Matrix(2, 2))
        img = Image.open(io.BytesIO(pix.tobytes("png")))
        items = reader.readtext(_preprocess(img), detail=1, paragraph=False)
        pages[str(pno + 1)] = {
            "objects": len(items),
            "texts": [t for _, t, _ in items],
            "confidences": [round(float(c), 3) for _, _, c in items],
            "lines": _boxes_to_lines(items),
        }
        log(f"  страница {pno + 1}: распознано объектов {len(items)}", "OCR")
    doc.close()
    data = {"file": os.path.basename(pdf_path), "pages": pages,
            "seconds": round(time.time() - t0, 1)}
    _save_ocr_cache(pdf_path, data)
    log(f"OCR завершён за {data['seconds']} с", "OCR")
    return data


def ocr_image(image_path: str) -> dict:
    """OCR скана/фото карточки (PNG, JPG, BMP, TIF, TIFF, WEBP) с кэшем."""
    from PIL import Image
    data = _load_ocr_cache(image_path)
    if data is not None:
        return data
    reader = ocr_reader()
    t0 = time.time()
    img = Image.open(image_path)
    items = reader.readtext(_preprocess(img), detail=1, paragraph=False)
    pages = {"1": {
        "objects": len(items),
        "texts": [t for _, t, _ in items],
        "confidences": [round(float(c), 3) for _, _, c in items],
        "lines": _boxes_to_lines(items),
    }}
    data = {"file": os.path.basename(image_path), "pages": pages,
            "seconds": round(time.time() - t0, 1)}
    _save_ocr_cache(image_path, data)
    log(f"OCR изображения завершён за {data['seconds']} с; объектов {len(items)}",
        "OCR")
    return data


def ocr_text(data: dict) -> str:
    """Собрать текст из результата OCR (страницы по порядку)."""
    lines = []
    for pno in sorted(data.get("pages", {}), key=lambda x: int(x)):
        lines.append(f"===== СТРАНИЦА {pno} =====")
        lines.extend(data["pages"][pno].get("lines", []))
    return "\n".join(lines)


def pdf_text(pdf_path: str, force_ocr: bool = False) -> tuple:
    """Текст карточки: (текст, источник). Слой PDF -> OCR."""
    import pymupdf
    if not force_ocr:
        doc = pymupdf.open(pdf_path)
        chunks = []
        for pno, page in enumerate(doc):
            t = page.get_text("text")
            if t.strip():
                chunks.append(f"===== СТРАНИЦА {pno + 1} =====")
                chunks.append(t.strip())
        doc.close()
        text = "\n".join(chunks)
        if len(text.strip()) > 200:
            log(f"использован ТЕКСТОВЫЙ СЛОЙ PDF: {len(text)} символов", "ТЕКСТ")
            return text, "текстовый слой PDF"
        log("текстового слоя нет (скан) — перехожу к OCR", "ТЕКСТ")
    data = ocr_pdf(pdf_path)
    text = ocr_text(data)
    log(f"OCR дал {len(text)} символов, строк "
        f"{sum(len(v.get('lines', [])) for v in data.get('pages', {}).values())}",
        "ТЕКСТ")
    return text, "OCR (EasyOCR ru+en)"


# ----------------------------------------------------------------------
# 2.1 ЧТЕНИЕ ДРУГИХ ФОРМАТОВ: DOCX, ODT, XLSX, TXT/MD, DOC/RTF, XLS, картинки
# ----------------------------------------------------------------------
def _xml_local(tag: str) -> str:
    """Имя XML-тега без пространства имён ('{ns}p' -> 'p')."""
    return str(tag).rsplit("}", 1)[-1]


def _docx_cell_text(cell) -> str:
    """Текст ячейки таблицы DOCX (абзацы + вложенные таблицы)."""
    parts = [p.text.strip() for p in cell.paragraphs if p.text.strip()]
    for nested in getattr(cell, "tables", []) or []:
        for row in nested.rows:
            parts.extend(t for t in (_docx_cell_text(c) for c in row.cells) if t)
    return " ".join(parts)


def docx_text(path: str) -> str:
    """Текст DOCX в порядке документа: абзацы, таблицы, колонтитулы."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = Document(path)
    chunks: list[str] = []

    def add_paragraph(par) -> None:
        t = par.text.strip()
        if t:
            chunks.append(t)

    def add_table(tbl) -> None:
        for row in tbl.rows:
            cells = []
            for cell in row.cells:
                val = _docx_cell_text(cell)
                # объединённые ячейки Word отдаёт повторно — не дублируем
                if val and (not cells or cells[-1] != val):
                    cells.append(val)
            if cells:
                chunks.append(" | ".join(cells))

    for child in doc.element.body.iterchildren():
        tag = _xml_local(child.tag)
        if tag == "p":
            add_paragraph(Paragraph(child, doc))
        elif tag == "tbl":
            add_table(Table(child, doc))
    for section in doc.sections:        # колонтитулы (название организации и т. п.)
        for part in (section.header, section.footer):
            if part is None:
                continue
            for par in part.paragraphs:
                t = par.text.strip()
                if t and t not in chunks:
                    chunks.append(t)
    return "\n".join(chunks)


def _odt_collect(el, out: list) -> None:
    """Собрать текст элемента ODT (text:p / text:h считаются абзацами)."""
    tag = _xml_local(el.tag)
    if tag in ("p", "h"):
        buf: list[str] = []
        if el.text:
            buf.append(el.text)
        for child in el:
            _odt_collect(child, buf)
            if child.tail:
                buf.append(child.tail)
        text = re.sub(r"[ \t]+", " ", "".join(buf)).strip()
        if text:
            out.append(text)
        return
    if tag == "tab":
        out.append("\t")
        return
    if tag in ("line-break", "br", "s"):
        out.append(" ")
        return
    if el.text:
        out.append(el.text)
    for child in el:
        _odt_collect(child, out)
        if child.tail:
            out.append(child.tail)


def odt_text(path: str) -> str:
    """Текст ODT (zip + content.xml) без внешних зависимостей."""
    import zipfile
    import xml.etree.ElementTree as ET
    with zipfile.ZipFile(path) as z:
        target = next((n for n in z.namelist()
                       if n.endswith("content.xml")), "")
        if not target:
            log("в ODT не найден content.xml", "ТЕКСТ")
            return ""
        root = ET.fromstring(z.read(target))
    chunks: list[str] = []

    def scan(el) -> None:
        if _xml_local(el.tag) in ("p", "h"):
            _odt_collect(el, chunks)
            return
        for child in el:
            scan(child)

    scan(root)
    return "\n".join(chunks)


def xlsx_text(path: str) -> str:
    """Текст XLSX (значения ячеек по листам) без openpyxl: zip + XML."""
    import zipfile
    import xml.etree.ElementTree as ET
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        shared: list[str] = []
        if "xl/sharedStrings.xml" in names:
            root = ET.fromstring(z.read("xl/sharedStrings.xml"))
            for si in root:
                if _xml_local(si.tag) != "si":
                    continue
                shared.append("".join(t.text or "" for t in si.iter()
                                      if _xml_local(t.tag) == "t"))
        sheet_re = re.compile(r"^xl/worksheets/sheet(\d+)\.xml$")
        sheets = sorted((int(m.group(1)), n) for n in names
                        for m in [sheet_re.match(n)] if m)
        chunks: list[str] = []
        for num, sheet in sheets:
            root = ET.fromstring(z.read(sheet))
            rows: list[str] = []
            for row in root.iter():
                if _xml_local(row.tag) != "row":
                    continue
                vals = []
                for cell in row:
                    if _xml_local(cell.tag) != "c":
                        continue
                    ctype = cell.get("t") or ""
                    raw, inline = "", ""
                    for child in cell:
                        ctag = _xml_local(child.tag)
                        if ctag == "v" and child.text is not None:
                            raw = child.text
                        elif ctag == "is":
                            inline = "".join(t.text or "" for t in child.iter()
                                             if _xml_local(t.tag) == "t")
                    val = inline or raw
                    if ctype == "s" and val.strip().isdigit():
                        idx = int(val)
                        val = shared[idx] if 0 <= idx < len(shared) else ""
                    val = re.sub(r"\s+", " ", val).strip()
                    if val:
                        vals.append(val)
                if vals:
                    rows.append(" | ".join(vals))
            if rows:
                chunks.append(f"===== ЛИСТ {num} =====")
                chunks.extend(rows)
    return "\n".join(chunks)


def plain_text(path: str) -> str:
    """Текст TXT/MD с автоопределением кодировки (UTF-8/UTF-16/CP1251)."""
    with open(path, "rb") as f:
        raw = f.read()
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16", errors="replace")
    for enc in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _remove_temp(tmp: str, folder: str) -> None:
    """Удалить временный файл и папку конвертации (ошибки не критичны)."""
    try:
        os.remove(tmp)
    except OSError:
        pass
    try:
        os.rmdir(folder)
    except OSError:
        pass


def _word_to_text(path: str) -> str:
    """DOC/RTF -> текст через установленный Microsoft Word (pywin32)."""
    import tempfile
    import pythoncom
    import win32com.client as win32
    tmp_dir = tempfile.mkdtemp(prefix="adapt_word_")
    tmp = os.path.join(tmp_dir,
                       os.path.splitext(os.path.basename(path))[0] + ".docx")
    pythoncom.CoInitialize()
    word = None
    try:
        word = win32.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        doc = word.Documents.Open(os.path.abspath(path), ReadOnly=True,
                                  AddToRecentFiles=False)
        try:
            try:
                doc.SaveAs2(tmp, FileFormat=16)      # wdFormatXMLDocument
            except AttributeError:                   # старые версии Word
                doc.SaveAs(tmp, FileFormat=16)
        finally:
            doc.Close(False)
    finally:
        if word is not None:
            word.Quit()
        pythoncom.CoUninitialize()
    try:
        return docx_text(tmp)
    finally:
        _remove_temp(tmp, tmp_dir)


def _xls_to_text(path: str) -> str:
    """XLS -> текст через установленный Microsoft Excel (pywin32)."""
    import tempfile
    import pythoncom
    import win32com.client as win32
    tmp_dir = tempfile.mkdtemp(prefix="adapt_xls_")
    tmp = os.path.join(tmp_dir,
                       os.path.splitext(os.path.basename(path))[0] + ".xlsx")
    pythoncom.CoInitialize()
    excel = None
    try:
        excel = win32.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        wb = excel.Workbooks.Open(os.path.abspath(path), ReadOnly=True)
        try:
            wb.SaveAs(tmp, FileFormat=51)            # xlOpenXMLWorkbook
        finally:
            wb.Close(False)
    finally:
        if excel is not None:
            excel.Quit()
        pythoncom.CoUninitialize()
    try:
        return xlsx_text(tmp)
    finally:
        _remove_temp(tmp, tmp_dir)


# Служебные группы RTF, текст которых в карточку не нужен
RTF_SKIP_GROUPS = (
    "fonttbl", "colortbl", "stylesheet", "info", "pict", "object", "header",
    "footer", "footnote", "annotation", "themedata", "datastore", "xmlnstbl",
    "listtable", "listoverridetable", "rsidtbl", "generator", "filetbl",
)


def rtf_text(path: str) -> str:
    """Резервный разбор RTF без Word (если Word не установлен)."""
    with open(path, "rb") as f:
        raw = f.read().decode("latin-1", errors="replace")
    out: list[str] = []
    stack: list[bool] = []
    skip = False
    i, n = 0, len(raw)
    while i < n:
        ch = raw[i]
        if ch == "{":
            stack.append(skip)
            m = re.match(r"\{\\([a-zA-Z]+)", raw[i:])
            if m and m.group(1) in RTF_SKIP_GROUPS:
                skip = True
            i += 1
            continue
        if ch == "}":
            skip = stack.pop() if stack else False
            i += 1
            continue
        if ch == "\\":
            m = re.match(r"\\([a-zA-Z]+)(-?\d+)? ?", raw[i:])
            if m:
                word_name, num = m.group(1), m.group(2)
                if word_name == "u" and num and not skip:
                    code = int(num)
                    out.append(chr(code + 65536 if code < 0 else code))
                    i += m.end()
                    if i < n and raw[i] == "?":     # fallback-символ \uN?
                        i += 1
                    continue
                elif word_name in ("par", "line", "sect", "page") and not skip:
                    out.append("\n")
                elif word_name == "tab" and not skip:
                    out.append("\t")
                i += m.end()
                continue
            if i + 1 < n and raw[i + 1] in "\\{}" and not skip:
                out.append(raw[i + 1])
            i += 2
            continue
        if not skip and ch not in "\r\n":
            out.append(ch)
        i += 1
    text = re.sub(r"[ \t]{2,}", " ", "".join(out))
    return "\n".join(l.strip() for l in text.splitlines() if l.strip())


def _extract_text_by_ext(card_path: str, ext: str, force_ocr: bool) -> tuple:
    """Прочитать карточку конкретного формата: (текст, источник)."""
    label = EXT_NAMES.get(ext, ext)
    if ext in PDF_EXTENSIONS:
        return pdf_text(card_path, force_ocr=force_ocr)
    if ext in IMAGE_EXTENSIONS:
        text = ocr_text(ocr_image(card_path))
        log(f"OCR изображения дал {len(text)} символов", "ТЕКСТ")
        return text, f"OCR изображения (EasyOCR ru+en, {label})"
    if ext in DOCX_EXTENSIONS:
        return docx_text(card_path), f"{label} (python-docx)"
    if ext in ODT_EXTENSIONS:
        return odt_text(card_path), f"{label} (content.xml)"
    if ext in PLAIN_EXTENSIONS:
        return plain_text(card_path), f"{label} (текстовый файл)"
    if ext in WORD_EXTENSIONS:
        try:
            return _word_to_text(card_path), f"{label} (Word)"
        except Exception as e:
            log(f"{label} не прочитан через Word ({type(e).__name__}: {e})",
                "ТЕКСТ")
            if ext == ".rtf":
                text = rtf_text(card_path)
                if text.strip():
                    log("использован встроенный разбор RTF", "ТЕКСТ")
                    return text, "RTF (встроенный разбор)"
        return "", f"{label}: не удалось прочитать"
    if ext in XLSX_EXTENSIONS:
        return xlsx_text(card_path), f"{label} (значения листов)"
    if ext in XLS_EXTENSIONS:
        try:
            return _xls_to_text(card_path), f"{label} (Excel)"
        except Exception as e:
            log(f"{label} не прочитан через Excel ({type(e).__name__}: {e})",
                "ТЕКСТ")
            return "", "XLS: не удалось прочитать (сохраните как .xlsx)"
    return "", f"{label or ext}: формат не поддерживается"


def extract_text(card_path: str, force_ocr: bool = False) -> tuple:
    """Текст карточки любого поддерживаемого формата: (текст, источник).

    PDF — текстовый слой/OCR; DOCX, ODT, XLSX, TXT/MD — чтение напрямую;
    DOC, RTF, XLS — через Word/Excel (pywin32); картинки — OCR EasyOCR.
    """
    name = os.path.basename(str(card_path))
    if not os.path.exists(card_path):
        log(f"ОШИБКА: файл не найден: {card_path}", "ТЕКСТ")
        return "", "файл не найден"
    ext = card_ext(card_path)
    if not is_supported_card(card_path):
        log(f"ОШИБКА: формат «{ext or name}» не поддерживается. Поддерживаются: "
            + ", ".join(dict.fromkeys(EXT_NAMES[e] for e in CARD_EXTENSIONS)),
            "ТЕКСТ")
        return "", f"формат «{ext or name}» не поддерживается"
    try:
        return _extract_text_by_ext(card_path, ext, force_ocr)
    except Exception as e:
        log(f"ОШИБКА чтения «{name}»: {type(e).__name__}: {e}", "ТЕКСТ")
        return "", f"ошибка чтения ({type(e).__name__})"


# ----------------------------------------------------------------------
# 3. ПОИСК ПРОФСТАНДАРТА / ЕТКС В ИНТЕРНЕТЕ
# ----------------------------------------------------------------------
REQ_KEYS = ("знать", "уметь", "трудовая функция", "трудовые действия",
            "профстандарт", "еткс", "квалификационн", "требования")


def search_web(query: str, max_results: int = 6) -> list:
    """Веб-поиск через ddgs. Возвращает [{title, href, body}]."""
    try:
        from ddgs import DDGS
    except Exception:
        try:
            from duckduckgo_search import DDGS
        except Exception as e:
            log(f"поиск недоступен (нет ddgs/duckduckgo_search): "
                f"{type(e).__name__}: {e}", "ПОИСК WEB")
            return []
    try:
        with DDGS() as ddgs:
            res = list(ddgs.text(query, max_results=max_results, region="ru-ru"))
        log(f"поиск «{query}»: результатов {len(res)}", "ПОИСК WEB")
        return [{"title": r.get("title", ""), "href": r.get("href", ""),
                 "body": r.get("body", "")} for r in res]
    except Exception as e:
        log(f"ошибка поиска «{query}»: {type(e).__name__}: {e}", "ПОИСК WEB")
        return []


def fetch_page_text(url: str, max_chars: int = 6000) -> str:
    """Скачать страницу и вытащить текст абзацев с требованиями."""
    try:
        import requests
        r = requests.get(url, timeout=20, headers={
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"})
        if r.status_code != 200:
            log(f"страница {url}: HTTP {r.status_code}", "ПОИСК WEB")
            return ""
        html = r.text
    except Exception as e:
        log(f"страница {url}: {type(e).__name__}: {e}", "ПОИСК WEB")
        return ""
    html = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", html)
    text = re.sub(r"(?s)<[^>]+>", "\n", html)
    text = (text.replace("&nbsp;", " ").replace("&amp;", "&")
                .replace("&quot;", '"').replace("&#39;", "'")
                .replace("&laquo;", "«").replace("&raquo;", "»"))
    lines = [re.sub(r"\s+", " ", l).strip() for l in text.splitlines()]
    keep = [l for l in lines if len(l) > 25 and any(k in l.lower()
                                                    for k in REQ_KEYS)]
    if not keep:
        keep = [l for l in lines if len(l) > 40][:80]
    out = "\n".join(keep)[:max_chars]
    log(f"со страницы {url[:70]}… взято {len(out)} символов требований", "ПОИСК WEB")
    return out


def profstandart_context(position: str, enabled: bool = True,
                         max_pages: int = 3) -> dict:
    """Найти профстандарт/ЕТКС по должности и вернуть требования текстом."""
    if not enabled:
        log("веб-поиск профстандарта отключён (--no-search)", "ПРОФСТАНДАРТ")
        return {"found": False, "sources": [], "text": ""}
    queries = [f"профстандарт {position} знать уметь трудовые действия",
               f"ЕТКС {position} квалификационные требования",
               f"должностная инструкция {position} требования к знаниям и умениям"]
    sources, parts = [], []
    for q in queries:
        for item in search_web(q, max_results=4):
            if item["href"] in [s["href"] for s in sources]:
                continue
            sources.append(item)
            if len(sources) >= max_pages:
                break
        if len(sources) >= max_pages:
            break
    for s in sources[:max_pages]:
        body = fetch_page_text(s["href"])
        if body:
            parts.append(f"### ИСТОЧНИК: {s['title']} ({s['href']})\n{body}")
    text = "\n\n".join(parts)
    log(f"профстандарт/ЕТКС: источников {len(sources)}, текста {len(text)} символов",
        "ПРОФСТАНДАРТ")
    return {"found": bool(text), "sources": sources, "text": text}


# ----------------------------------------------------------------------
# 4. ПРОФИЛЬ ДОЛЖНОСТИ (JSON) — рукописные карточки распознаются плохо,
#    поэтому поддерживается сайдкар-файл с полями профиля
# ----------------------------------------------------------------------
def load_sidecar_profile(card_path: str) -> dict:
    """Найти JSON-профиль рядом с карточкой (или в _cache/_profiles)."""
    stem = os.path.splitext(os.path.basename(card_path))[0]
    candidates = [
        os.path.join(os.path.dirname(card_path), stem + ".json"),
        os.path.join(CACHE_FOLDER, stem + "_profile.json"),
        os.path.join(OUT_FOLDER, "_profiles", stem + ".json"),
    ]
    for path in candidates:
        if os.path.exists(path):
            try:
                with open(path, encoding="utf-8") as f:
                    data = json.load(f)
                log(f"найден JSON-профиль: {path}", "ПРОФИЛЬ")
                return data
            except Exception as e:
                log(f"JSON-профиль {path} не прочитан: {type(e).__name__}: {e}",
                    "ПРОФИЛЬ")
    return {}


# ----------------------------------------------------------------------
# 5. АНАЛИЗ ПРОФИЛЯ ДОЛЖНОСТИ ЛОКАЛЬНОЙ LLM -> JSON
# ----------------------------------------------------------------------
SYSTEM_PROMPT = """Ты — эксперт по адаптации персонала и разработке планов развития
с большим опытом. Ты анализируешь профиль должности работника (ПД) и формируешь
индивидуальный план адаптации сроком до 3 месяцев по рабочей инструкции
«Адаптация персонала» (РИ 7.2-01-2024, Приложение Б).

ПРАВИЛА:
1. Из профиля извлекаются поля: Ф.И.О., Подразделение, Должность, Руководитель,
   Наставник, Менеджер программ обучения, категория должности, срок адаптации.
   Если поле не читается — оставь пустую строку ("").
2. Профиль содержит блоки «4.1 Теоретические знания» и «4.2 Практические навыки»
   с оценками от 0 до 5. НОРМАТИВ — 4 балла. Позиции с оценкой меньше 4 (или без
   оценки) считаются ПРОБЕЛАМИ и попадают в "пробелы" и в индивидуальный блок.
3. Если в профиле мало информации о знаниях и умениях, требования бери из
   переданного текста профстандарта/ЕТКС/должностной инструкции (блок
   ПРОФСТАНДАРТ) и укажи это в "источник_требований".
4. Мероприятия УКРУПНЯЙ по функциональным признакам (3-6 пунктов на блок),
   без микро-шагов. Порядок — от ранних сроков к поздним, логичный.
5. Даты бери РОВНО те, что переданы в блоке ОКНА ДАТ (не пересчитывай).
6. «Общий» блок — организация, охрана труда, правила внутреннего распорядка,
   знакомство с подразделением и командой. «Индивидуальный» — рабочее место,
   оборудование, документация, освоение трудовых функций (самостоятельно или с
   наставником) по выявленным пробелам. «Итоговое задание» — конкретное
   практическое задание или тест на готовность к самостоятельной работе.
7. В поле "участники" ВСЕГДА указывай ДОЛЖНОСТЬ и ФАМИЛИЮ С ИНИЦИАЛАМИ
   работника в формате «Должность — Фамилия И.О.» (несколько участников —
   через «; »). Адаптируемый работник указывается в КАЖДОМ мероприятии своей
   должностью. Пример: «Инженер по качеству 1 категории — Шварева Ю.С.;
   Руководитель структурного подразделения — Галаган И.Г.».
8. ВСЕ поля плана обязательны, кроме подписи адаптируемого работника. Если
   данные по полю получить не удалось — оставь ПУСТУЮ строку (""), система
   сама выведет это поле красным текстом «не удалось получить данные».

ОТВЕТ — СТРОГО JSON (без markdown) по схеме:
{
  "фио": "", "подразделение": "", "должность": "", "руководитель": "",
  "наставник": "", "менеджер_программ_обучения": "", "категория_должности": "",
  "не_требуется": false, "причина_исключения": "",
  "пробелы": [{"позиция": "", "оценка": 0, "норматив": 4, "комментарий": ""}],
  "источник_требований": "",
  "общий_блок": [{"мероприятие": "", "участники": "", "проверка_знаний": ""}],
  "индивидуальный_блок": [{"мероприятие": "", "участники": "", "проверка_знаний": ""}],
  "итоговое_задание": {"задание": "", "критерии": "", "участники": ""}
}"""


def _repair_json(text: str) -> str:
    """Достроить незакрытые кавычки/скобки в ОБРЕЗАННОМ ответе модели."""
    out, stack, in_str, esc = [], [], False, False
    for ch in text:
        if in_str:
            if esc:
                esc = False
            elif ch == "\\":
                esc = True
            elif ch == '"':
                in_str = False
        else:
            if ch == '"':
                in_str = True
            elif ch in "{[":
                stack.append(ch)
            elif ch in "}]" and stack:
                stack.pop()
        out.append(ch)
    res = "".join(out)
    if in_str:
        res += '"'
    for ch in reversed(stack):
        res += "}" if ch == "{" else "]"
    return res


def _parse_json(text: str) -> dict:
    """Разобрать JSON: как есть -> подстрока {...} -> восстановление скобок."""
    if not text:
        return {}
    for candidate in (text, (re.search(r"(?s)\{.*\}", text) or [None])[0]):
        if not candidate:
            continue
        try:
            return json.loads(candidate)
        except Exception:
            try:
                data = json.loads(_repair_json(candidate))
                if isinstance(data, dict) and data:
                    return data
            except Exception:
                continue
    return {}


def llm_json(messages: list, model: str = MODEL, tag: str = "LLM",
             retry: bool = True) -> dict:
    """Запрос к Ollama с принудительным JSON, разбором и повтором при обрыве."""
    import ollama
    kwargs = {"model": model, "format": "json", "messages": messages,
              "options": {"temperature": LLM_TEMPERATURE,
                          "num_ctx": LLM_NUM_CTX},
              "think": LLM_THINK}
    if OLLAMA_HOST:
        kwargs["host"] = OLLAMA_HOST
    r, last = None, None
    while r is None:               # постепенно убираем неподдерживаемые параметры
        try:
            r = ollama.chat(**kwargs)
        except Exception as e:
            last = e
            if "think" in kwargs:  # модель/клиент без поддержки think
                kwargs.pop("think")
                continue
            if "host" in kwargs:   # старые версии клиента без host
                kwargs.pop("host")
                continue
            break
    if r is None:
        log(f"ОШИБКА вызова модели {model}: {type(last).__name__}: {last}", tag)
        return {}
    content = (r.get("message") or {}).get("content", "") or ""
    data = _parse_json(content)
    if data:
        return data
    log(f"JSON не разобран (длина ответа {len(content)} символов)", tag)
    if retry:
        log("повторный запрос: сокращённый ввод и требование компактного JSON", tag)
        short = []
        for m in messages:
            c = str(m.get("content", ""))
            if m.get("role") == "user":
                c = (c[:6000] + "\n\nВАЖНО: верни СТРОГО компактный JSON и "
                     "заверши его полностью; не более 4 мероприятий в каждом "
                     "блоке, текст мероприятий — до 120 символов.")
            short.append({"role": m.get("role", "user"), "content": c})
        return llm_json(short, model=model, tag=tag, retry=False)
    log(f"модель вернула не JSON: {content[:200]!r}", tag)
    return {}


def analyze_profile(text: str, sidecar: dict, prof: dict, start: datetime,
                    end: datetime, model: str = MODEL) -> dict:
    """LLM-анализ профиля и формирование структуры плана (окна дат считает код)."""
    gen_win = (start, start + timedelta(days=GENERAL_DAYS))
    fin_win = (end - timedelta(days=FINAL_DAYS), end)
    ind_win = (gen_win[1] + timedelta(days=1), fin_win[0] - timedelta(days=1))
    if ind_win[1] < ind_win[0]:
        ind_win = (gen_win[1], fin_win[0])
    parts = [
        "ОКНА ДАТ (используй ровно эти значения, не пересчитывай):\n"
        f"  срок адаптации: с {start.strftime('%d.%m.%Y')} "
        f"по {end.strftime('%d.%m.%Y')};\n"
        f"  общий блок: {date_range(*gen_win)};\n"
        f"  индивидуальный блок: {date_range(*ind_win)};\n"
        f"  итоговое задание: {date_range(*fin_win)};\n"
        f"  дата составления: {datetime.now().strftime('%d.%m.%Y')}",
        "\n=== ТЕКСТ ПРОФИЛЯ ДОЛЖНОСТИ (OCR/текстовый слой) ===\n" + text[:9000],
    ]
    if sidecar:
        parts.append("\n=== ПОЛЯ ПРОФИЛЯ (заполнены человеком, приоритет выше OCR) ===\n"
                     + json.dumps(sidecar, ensure_ascii=False, indent=1)[:4000])
    if prof.get("found"):
        parts.append("\n=== ПРОФСТАНДАРТ/ЕТКС (требования из интернета) ===\n"
                     + prof["text"][:4000])
        parts.append("Источники: " + "; ".join(
            s["href"] for s in prof.get("sources", [])[:3]))
    parts.append("\nСформируй план адаптации строго по схеме JSON.")
    log(f"запрос к LLM {model}: профиль {len(text)} симв., профстандарт "
        f"{len(prof.get('text', ''))} симв.", "LLM")
    t0 = time.time()
    data = llm_json([{"role": "system", "content": SYSTEM_PROMPT},
                     {"role": "user", "content": "\n".join(parts)}], model=model)
    log(f"ответ LLM получен за {time.time() - t0:.1f} с; полей в JSON: {len(data)}",
        "LLM")
    if not data:
        return {}
    for key in ("фио", "подразделение", "должность", "руководитель", "наставник",
                "менеджер_программ_обучения", "категория_должности",
                "источник_требований", "причина_исключения"):
        if not isinstance(data.get(key), str):
            data[key] = "" if data.get(key) is None else str(data.get(key))
    data["срок_адаптации"] = (f"с {start.strftime('%d.%m.%Y')} "
                              f"по {end.strftime('%d.%m.%Y')}")
    data["дата_составления"] = datetime.now().strftime("%d.%m.%Y")
    for key in ("пробелы", "общий_блок", "индивидуальный_блок"):
        if not isinstance(data.get(key), list):
            data[key] = []
    if not isinstance(data.get("итоговое_задание"), dict):
        data["итоговое_задание"] = {
            "задание": str(data.get("итоговое_задание") or ""), "критерии": ""}
    data["_окна"] = {"общий": date_range(*gen_win),
                     "индивидуальный": date_range(*ind_win),
                     "итоговое": date_range(*fin_win)}
    data["_окна_dt"] = {"общий": gen_win, "индивидуальный": ind_win,
                        "итоговое": fin_win}
    return data


def apply_sidecar(data: dict, sidecar: dict) -> dict:
    """Поля из JSON-профиля (заполнен человеком) важнее ответа OCR-модели."""
    if not sidecar:
        return data
    mapping = {
        "фио": ("фио", "ФИО", "fio", "employee"),
        "подразделение": ("подразделение", "department", "структурное_подразделение"),
        "должность": ("должность", "position", "профессия"),
        "руководитель": ("руководитель", "manager"),
        "наставник": ("наставник", "mentor"),
        "менеджер_программ_обучения": ("менеджер_программ_обучения",
                                       "менеджер программ обучения", "hr"),
        "генеральный_директор": ("генеральный_директор", "генеральный директор"),
        "руководитель_удц": ("руководитель_удц", "руководитель УДЦ"),
        "категория_должности": ("категория_должности", "категория"),
        "срок_адаптации": ("срок_адаптации", "период_адаптации", "срок"),
        "дата_составления": ("дата_составления", "дата"),
    }
    for dst, srcs in mapping.items():
        for src in srcs:
            val = sidecar.get(src)
            if isinstance(val, str) and val.strip():
                data[dst] = val.strip()
                break
    for key in ("теоретические_знания", "практические_навыки", "пробелы"):
        if isinstance(sidecar.get(key), list) and not data.get(key):
            data[key] = sidecar[key]
    return data


# ----------------------------------------------------------------------
# 6. ФОРМИРОВАНИЕ WORD-ПЛАНА ПО ШАБЛОНУ
# ----------------------------------------------------------------------
def fill_cell(cell, lines) -> None:
    """Заполнить ячейку строками, сохранив формат шаблона (строка -> 1 строка).

    Строка может быть строкой либо парой (текст, красный_флаг) — красным
    помечаются поля, данные по которым получить не удалось.
    """
    from docx.text.paragraph import Paragraph
    if isinstance(lines, str) or lines is None:
        lines = [lines]
    prepared = []
    for x in (lines or [""]):
        if isinstance(x, (tuple, list)):
            text, red = (x[0], bool(x[1])) if len(x) > 1 else (x[0], False)
        else:
            text, red = x, False
        prepared.append(("" if text is None else str(text), red))
    p0 = cell.paragraphs[0]
    proto = copy.deepcopy(p0._p)              # образец абзаца (pPr + rPr)
    for extra in cell.paragraphs[1:]:         # убрать лишние абзацы
        extra._element.getparent().remove(extra._element)
    for r in list(p0.runs):                   # очистить первый абзац
        r._element.getparent().remove(r._element)
    run = p0.add_run(prepared[0][0])
    if prepared[0][1]:
        run.font.color.rgb = missing_color()
    anchor = p0._p
    for text, red in prepared[1:]:
        new_p = copy.deepcopy(proto)
        for child in list(new_p):
            if child.tag.endswith("}r"):      # выкинуть текстовые run'ы образца
                new_p.remove(child)
        anchor.addnext(new_p)
        anchor = new_p
        run = Paragraph(new_p, cell).add_run(text)
        if red:
            run.font.color.rgb = missing_color()


def fill_paragraph(p, text: str) -> None:
    """Заменить текст абзаца, сохранив формат первого run'а."""
    proto = copy.deepcopy(p.runs[0]._element) if p.runs else None
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    if proto is not None:
        for child in list(proto):
            if child.tag.endswith("}t"):
                proto.remove(child)
        p._p.append(proto)
        p.runs[-1].text = text
    else:
        p.add_run(text)


def set_default_tab_stop(doc, twips: int = 14175) -> None:
    """Убрать дефолтные табуляции (25 см), чтобы они не ломали выравнивание.

    Иначе Word сам вставляет таб-стопы каждые 1,27 см, и короткие подписи
    («Работник») «уезжают» влево от общих колонок.
    """
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    try:
        settings = doc.settings.element
    except Exception:
        settings = doc.part.settings.element
    el = settings.find(qn("w:defaultTabStop"))
    if el is None:
        el = OxmlElement("w:defaultTabStop")
        settings.append(el)
    el.set(qn("w:val"), str(int(twips)))


def insert_empty_paragraph_after(p):
    """Вставить пустой абзац сразу после указанного (сохранив формат)."""
    from docx.text.paragraph import Paragraph
    new_el = copy.deepcopy(p._p)
    for child in list(new_el):
        if child.tag.endswith("}r"):
            new_el.remove(child)
    p._p.addnext(new_el)
    return Paragraph(new_el, p._parent)


def normalize_fonts(doc, size_pt: int = FONT_SIZE_PT) -> None:
    """Единый шрифт плана: Times New Roman нужного кегля.

    * шапка (таблица 1) — все ячейки;
    * таблица мероприятий (таблица 2) — все строки, КРОМЕ строки «оглавления»
      (строки с названиями столбцов);
    * нижняя часть под таблицами — где было 12 пт (размеры 9 и 10 пт не трогаем).
    """
    from docx.shared import Pt
    from docx.text.paragraph import Paragraph

    def fix(paragraphs):
        for p in paragraphs:
            for r in p.runs:
                r.font.name = FONT_NAME
                r.font.size = Pt(size_pt)

    for table in doc.tables[:1]:                   # шапка — целиком
        for row in table.rows:
            for cell in row.cells:
                fix(cell.paragraphs)
    for table in doc.tables[1:]:                   # мероприятия — без «оглавления»
        for ri, row in enumerate(table.rows):
            if ri == 0:
                continue
            for cell in row.cells:
                fix(cell.paragraphs)
    # нижняя часть (абзацы после последней таблицы): 12 пт -> size_pt
    seen_table = False
    for child in doc.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "tbl":
            seen_table = True
        elif tag == "p" and seen_table:
            for r in Paragraph(child, None).runs:
                if r.font.size is not None and r.font.size.pt == 12:
                    r.font.name = FONT_NAME
                    r.font.size = Pt(size_pt)
                    r.font.color.rgb = font_color()   # «авто» -> явный чёрный


def ensure_sign_line(p) -> bool:
    """ГАРАНТИЯ: в строке подписи всегда есть линия «_______________».

    Если по какой-то причине линии в абзаце нет (в т.ч. в строке наставника,
    когда наставник не назначен), она вставляется сразу после первой табуляции
    (перед Ф.И.О.). Возвращает True, если линию пришлось добавить.
    """
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    if SIGN_LINE in p.text:
        return False
    for r in p.runs:
        for node in list(r._element):
            if node.tag == qn("w:tab"):
                t = OxmlElement("w:t")
                t.set(qn("xml:space"), "preserve")
                t.text = SIGN_LINE
                node.addnext(t)
                log("строка подписи: линии не было — добавлена", "WORD")
                return True
    log("ВНИМАНИЕ: в строке подписи нет табуляции — линия не добавлена", "WORD")
    return False


def fill_signature(p, label: str, name, date: str, avail_cm: float = 16.75,
                   note: str = "") -> None:
    """Строка подписи в формате «Образец.docx»:

        <Должность>  →  _______________  →  Фамилия И.О.  →  дата
        [(для …)]           Подпись (по центру линии)

    Шрифт Times New Roman 11 пт (слово «Подпись» — 10 пт, пояснение «(для …)» —
    9 пт); колонки выравниваются явными таб-стопами (SIGN_TAB_FRACS), поэтому
    во всех строках линия для подписи, фамилия и дата стоят на одной горизонтали.
    Пояснение ``note`` печатается слева на строке со словом «Подпись», при этом
    само слово «Подпись» не сдвигается (оно центрируется своим таб-стопом).
    Ф.И.О. выводится красным, если данные получить не удалось.
    """
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Cm, Pt
    from docx.text.paragraph import Paragraph
    proto_run = copy.deepcopy(p.runs[0]._element) if p.runs else None
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    stops = p.paragraph_format.tab_stops
    stops.clear_all()
    for frac in SIGN_TAB_FRACS:
        stops.add_tab_stop(Cm(round(avail_cm * frac, 2)), WD_TAB_ALIGNMENT.LEFT)
    # центр линии подписи — по нему центрируется слово «Подпись»
    mid_cm = round(avail_cm * SIGN_TAB_FRACS[0] + SIGN_LINE_WIDTH_CM / 2.0, 2)
    stops.add_tab_stop(Cm(mid_cm), WD_TAB_ALIGNMENT.CENTER)
    proto_p = copy.deepcopy(p._p)          # pPr с таб-стопами, без ранов

    def add(text: str, red: bool = False, size_pt: int = FONT_SIZE_PT):
        if proto_run is not None:
            el = copy.deepcopy(proto_run)
            for child in list(el):
                if child.tag.endswith("}t"):
                    el.remove(child)
            p._p.append(el)
            run = p.runs[-1]
        else:
            run = p.add_run()
        run.text = text
        run.font.name = FONT_NAME
        run.font.size = Pt(size_pt)
        run.font.color.rgb = font_color()      # явный чёрный (в шаблоне — «авто»)
        if red:
            run.font.color.rgb = missing_color()
        return run

    # строка 1: должность | линия для подписи | Фамилия И.О. | дата.
    # Линия подписи — ОТДЕЛЬНЫМ раном со своим явным оформлением, чтобы
    # отображалась одинаково во всех строках и просмотрщиках.
    add(f"{label}\t")
    add(f"{SIGN_LINE}\t")
    text, red = field_line(name)
    add(text, red)
    add(f"\t{date}")
    # строка 2: [пояснение «(для …)»] + слово «Подпись» по центру линии (10 пт)
    p._p.addnext(proto_p)
    hint = Paragraph(proto_p, p._parent)
    if note:
        rn = hint.add_run(note)
        rn.font.name = FONT_NAME
        rn.font.size = Pt(SIGN_NOTE_PT)
        rn.font.color.rgb = font_color()
    run = hint.add_run(f"\t\t{SIGN_WORD}")
    run.font.name = FONT_NAME
    run.font.size = Pt(SIGN_WORD_PT)
    run.font.color.rgb = font_color()


def clone_row_after(table, src_row, values: list):
    """Скопировать строку таблицы (с форматом) и заполнить значениями."""
    from docx.table import _Row
    new_tr = copy.deepcopy(src_row._tr)
    src_row._tr.addnext(new_tr)
    new_row = _Row(new_tr, table)
    for cell, lines in zip(new_row.cells, values):
        fill_cell(cell, lines)
    return new_row


def block_rows(block: list, dates: list, numbers: str, check_default: str) -> list:
    """Строки таблицы блока: (№, мероприятие, участники, даты, подпись, проверка).

    Все поля обязательны, КРОМЕ столбца «Подпись адаптируемого работника»:
    он всегда остаётся пустым. Незаполненные поля помечаются красным.
    """
    rows = []
    for i, item in enumerate(block):
        if not isinstance(item, dict):
            item = {"мероприятие": str(item)}
        date = dates[i] if i < len(dates) else (dates[-1] if dates else "")
        rows.append([
            [(f"{numbers}.{i + 1}", False)],
            [field_line(item.get("мероприятие"))],
            [field_line(item.get("участники"))],
            [field_line(date)],
            [("", False)],                 # «Подпись адаптируемого работника» — пусто
            [field_line(item.get("проверка_знаний") or check_default)],
        ])
    return rows


# ----------------------------------------------------------------------
# 5.1 УЧАСТНИКИ: всегда «Должность — Фамилия И.О.»
# ----------------------------------------------------------------------
PARTICIPANT_ROLES = (
    ("Адаптируемый работник", ("адаптируем", "работник", "сотрудник", "новичок",
                               "стажер", "стажёр", "принят")),
    ("Руководитель структурного подразделения",
     ("руководител", "начальник", "заведующ", "мастер", "директор")),
    ("Менеджер программ обучения",
     ("менеджер программ", "менеджер по адаптации", "hr", "кадр", "удц")),
    ("Наставник", ("наставник", "mentor", "куратор")),
    ("Специалист по охране труда", ("охране труда", "охрана труда", "промбезопасн",
                                    "пожарн", "экологи", "от и пб")),
    ("Коллектив подразделения", ("коллектив", "команда", "подразделение", "отдел")),
)


def role_fio(role: str, data: dict) -> str:
    """Ф.И.О. по роли (пустая строка, если неизвестно/не назначено)."""
    val = str({
        "Адаптируемый работник": data.get("фио"),
        "Руководитель структурного подразделения": data.get("руководитель"),
        "Менеджер программ обучения": (data.get("менеджер_программ_обучения")
                                       or PROGRAM_MANAGER),
        "Наставник": data.get("наставник"),
    }.get(role) or "").strip()
    if norm(val) in ("", "нет", "не назначен", "не назначено", "-", "no"):
        return ""
    return val


def participant_text(role: str, data: dict) -> str:
    """«Должность — Фамилия И.О.» (Ф.И.О. — только если оно известно).

    Адаптируемый работник выводится его ДОЛЖНОСТЬЮ (поле «должность»), а не
    обобщённой ролью: «Инженер по качеству 1 категории — Шварева Ю.С.».
    """
    short = short_fio(role_fio(role, data))
    if role == "Коллектив подразделения":
        dep = str(data.get("подразделение") or "").strip()
        return f"Коллектив подразделения — {dep}" if dep else role
    if role == "Наставник" and not short:
        return "Наставник (не назначен)"
    if role == "Адаптируемый работник":
        pos = str(data.get("должность") or "").strip() or "Адаптируемый работник"
        return f"{pos} — {short}" if short else pos
    return f"{role} — {short}" if short else role


def participant_role(text: str, data: dict) -> str:
    """Роль участника по фрагменту текста: сначала ищем известные Ф.И.О."""
    t = norm(text)
    if not t:
        return ""
    for role in ("Адаптируемый работник", "Руководитель структурного подразделения",
                 "Менеджер программ обучения", "Наставник"):
        fio = role_fio(role, data)
        if not fio:
            continue
        variants = {norm(fio), norm(short_fio(fio))}
        parts = [p for p in re.split(r"\s+", fio) if p]
        if parts:                      # совпадение по одной фамилии
            variants.add(norm(parts[0]))
        if any(v and v in t for v in variants):
            return role
    for role, keys in PARTICIPANT_ROLES:
        if any(k in t for k in keys):
            return role
    return ""


def _split_participants(raw: str) -> list:
    return [p.strip(" .;,") for p in re.split(r"[;,/]|\s+и\s+", str(raw or ""))
            if p.strip(" .;,")]


def normalize_participants(data: dict) -> int:
    """Привести «Участники» во всех мероприятиях к виду «Должность — Ф.И.О.»."""
    employee = participant_text("Адаптируемый работник", data)
    fixed = 0
    for key in ("общий_блок", "индивидуальный_блок"):
        for item in data.get(key, []):
            if not isinstance(item, dict):
                continue
            raw = str(item.get("участники") or "").strip()
            out, seen = [], set()
            for part in _split_participants(raw):
                role = participant_role(part, data)
                txt = participant_text(role, data) if role else part
                if txt and norm(txt) not in seen:
                    seen.add(norm(txt))
                    out.append(txt)
            if employee and norm(employee) not in seen:
                out.insert(0, employee)
            new = "; ".join(out)
            if new != raw:
                fixed += 1
            item["участники"] = new
    task = data.get("итоговое_задание")
    if isinstance(task, dict):
        out, seen = [], set()
        for part in _split_participants(task.get("участники")):
            role = participant_role(part, data)
            txt = participant_text(role, data) if role else part
            if txt and norm(txt) not in seen:
                seen.add(norm(txt))
                out.append(txt)
        for role in ("Адаптируемый работник",
                     "Руководитель структурного подразделения", "Наставник"):
            txt = participant_text(role, data)
            if "не назначен" in txt or norm(txt) in seen:
                continue
            out.append(txt)
            seen.add(norm(txt))
        task["участники"] = "; ".join(out)
    return fixed


def build_plan_docx(data: dict, out_path: str) -> bool:
    """Заполнить шаблон плана адаптации (Приложение Б РИ) и сохранить."""
    from docx import Document
    from docx.shared import Emu
    if not os.path.exists(TEMPLATE):
        log(f"ОШИБКА: шаблон не найден: {TEMPLATE}", "WORD")
        return False
    doc = Document(TEMPLATE)
    if len(doc.tables) < 2:
        log("ОШИБКА: в шаблоне меньше двух таблиц — структура не распознана", "WORD")
        return False
    # --- таблица 1: шапка. Полное Ф.И.О. — ТОЛЬКО у адаптируемого работника;
    #     остальные Ф.И.О. — в формате «Фамилия И.О.». ---
    set_default_tab_stop(doc)
    t1 = doc.tables[0]
    header_fields = (
        ("фио", 0, False),
        ("подразделение", 1, False),
        ("должность", 2, False),
        ("руководитель", 3, True),
        ("менеджер_программ_обучения", 4, True),
        ("наставник", 5, True),
    )
    for key, row_idx, short in header_fields:
        val = data.get(key) or ""
        if key == "наставник" and not str(val).strip():
            val = "нет"          # наставник не назначен — допустимое значение
        if key == "менеджер_программ_обучения" and not str(val).strip():
            val = PROGRAM_MANAGER
        if short:
            val = short_fio(val)
        fill_cell(t1.rows[row_idx].cells[2], [field_line(val)])
    fill_cell(t1.rows[6].cells[1], [field_line(data.get("срок_адаптации"))])
    fill_cell(t1.rows[6].cells[3], [field_line(data.get("дата_составления"))])
    log(f"таблица 1: {data.get('фио')} | {data.get('должность')} | "
        f"{data.get('подразделение')} | рук.: {data.get('руководитель')}", "WORD")
    # --- таблица 2: три блока мероприятий ---
    t2 = doc.tables[1]
    win, win_dt = data.get("_окна", {}), data.get("_окна_dt", {})
    gen_rows = block_rows(data.get("общий_блок", []),
                          [win.get("общий", "")] * 20, "1",
                          "Проверка знаний не требуется")
    ind_items = data.get("индивидуальный_блок", [])
    if ind_items and win_dt.get("индивидуальный"):
        ind_windows = [date_range(a, b) for a, b in split_period(
            win_dt["индивидуальный"][0], win_dt["индивидуальный"][1],
            parts=len(ind_items))]
    else:
        ind_windows = [win.get("индивидуальный", "")] * len(ind_items)
    ind_rows = block_rows(ind_items, ind_windows, "2",
                          "Проверка знаний не требуется")
    task = data.get("итоговое_задание", {}) or {}
    fin_rows = [[
        [("3.", False)],
        [("Итоговое индивидуальное практическое задание:", False),
         field_line(task.get("задание"))],
        [field_line(task.get("участники")
                    or participant_text("Адаптируемый работник", data))],
        [field_line(win.get("итоговое"))],
        [("", False)],                 # «Подпись адаптируемого работника» — пусто
        [field_line(str(task.get("критерии", "")).strip()
                    or "Оценка практических умений и навыков: демонстрация "
                       "освоенных компетенций по итогам выполнения "
                       "практических заданий")],
    ]]
    for anchor_row, rows in ((t2.rows[2], gen_rows), (t2.rows[4], ind_rows),
                             (t2.rows[5], fin_rows)):
        if not rows:
            continue
        for cell, lines in zip(anchor_row.cells, rows[0]):
            fill_cell(cell, lines)
        prev = anchor_row
        for r in rows[1:]:
            prev = clone_row_after(t2, prev, r)
    log(f"таблица 2: общий блок — {len(gen_rows)} строк, индивидуальный — "
        f"{len(ind_rows)}, итоговое задание — {len(fin_rows)}", "WORD")
    # --- подписи и утверждение (все поля обязательны; нет данных — красным).
    #     Под таблицами имена — только «Фамилия И.О.»; полное Ф.И.О. остаётся
    #     в шапке (таблица 1). Колонки выравниваются явными таб-стопами.
    #     Пояснения «(для …)» печатаются на строке со словом «Подпись». ---
    dates = data.get("дата_составления", "")
    sec = doc.sections[0]
    avail_cm = Emu(sec.page_width - sec.left_margin - sec.right_margin).cm
    sign_map = (
        ("подразделения", "подразделения",
         short_fio(data.get("руководитель")), ""),
        ("Наставник", "Наставник (при назначении)",
         short_fio(data.get("наставник") or "нет"), ""),
        ("Менеджер программ обучения", "Менеджер программ обучения",
         short_fio(data.get("менеджер_программ_обучения") or PROGRAM_MANAGER), ""),
        ("Генеральный директор", "Генеральный директор",
         short_fio(data.get("генеральный_директор") or GENERAL_DIRECTOR), NOTE_GD),
        ("Руководитель УДЦ", "Руководитель УДЦ",
         short_fio(data.get("руководитель_удц") or UDC_HEAD), NOTE_UDC),
        ("Работник", "Работник", short_fio(data.get("фио")), ""),
    )
    # отдельные строки-пояснения «(для …)» больше не нужны — они на строке «Подпись»
    for p in list(doc.paragraphs):
        if p.text.strip().startswith("(для "):
            p._p.getparent().remove(p._p)
    for p in doc.paragraphs:
        txt = p.text.strip()
        if "Подпись" not in txt:
            continue
        for key, label, name, note in sign_map:
            if txt.startswith(key):
                fill_signature(p, label, name, dates, avail_cm=avail_cm, note=note)
                ensure_sign_line(p)      # линия подписи всегда на месте
                break
    # пустая строка после блока «Генеральный директор» (как в «Образец.docx»)
    for p in doc.paragraphs:
        if NOTE_GD in p.text:
            insert_empty_paragraph_after(p)
            break
    normalize_fonts(doc)
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    doc.save(out_path)
    log(f"план сохранён: {out_path}", "WORD")
    return True


# ----------------------------------------------------------------------
# 6.1 КОНТРОЛЬ ЗАПОЛНЕНИЯ ПОЛЕЙ ПЛАНА
# ----------------------------------------------------------------------
PLAN_FIELDS = (
    ("Ф.И.О. адаптируемого работника", lambda d: d.get("фио")),
    ("Подразделение", lambda d: d.get("подразделение")),
    ("Должность", lambda d: d.get("должность")),
    ("Ф.И.О. руководителя структурного подразделения",
     lambda d: d.get("руководитель")),
    ("Ф.И.О. менеджера программ обучения",
     lambda d: d.get("менеджер_программ_обучения") or PROGRAM_MANAGER),
    ("Срок адаптации", lambda d: d.get("срок_адаптации")),
    ("Дата составления", lambda d: d.get("дата_составления")),
    ("Ф.И.О. генерального директора",
     lambda d: d.get("генеральный_директор") or GENERAL_DIRECTOR),
    ("Ф.И.О. руководителя УДЦ",
     lambda d: d.get("руководитель_удц") or UDC_HEAD),
)


def audit_plan(data: dict) -> list:
    """Проверить обязательные поля плана (кроме «Подпись адаптируемого работника»).

    Возвращает список незаполненных полей: они выводятся в документе КРАСНЫМ
    «не удалось получить данные». Список пишется в JSON результата (_не_заполнено).
    """
    missing = [name for name, get in PLAN_FIELDS if is_missing(get(data))]
    for key in ("общий_блок", "индивидуальный_блок"):
        for i, item in enumerate(data.get(key, []) or [], 1):
            if not isinstance(item, dict):
                continue
            checks = (
                ("мероприятие", item.get("мероприятие")),
                ("участники", item.get("участники")),
                ("проверка_знаний", item.get("проверка_знаний")
                 or "Проверка знаний не требуется"),
            )
            for fname, fval in checks:
                if is_missing(fval):
                    missing.append(f"{key}[{i}].{fname}")
    task = data.get("итоговое_задание")
    if isinstance(task, dict):
        checks = (
            ("задание", task.get("задание")),
            ("участники", task.get("участники")
             or participant_text("Адаптируемый работник", data)),
            ("критерии", str(task.get("критерии", "")).strip()
             or "Оценка практических умений и навыков: демонстрация освоенных "
                "компетенций по итогам выполнения практических заданий"),
        )
        for fname, fval in checks:
            if is_missing(fval):
                missing.append(f"итоговое_задание.{fname}")
    if missing:
        log(f"ПОЛЯ БЕЗ ДАННЫХ ({len(missing)}) — выведены красным "
            f"«{MISSING_TEXT}»: " + "; ".join(missing), "КОНТРОЛЬ")
    else:
        log("все поля плана заполнены (кроме столбца «Подпись адаптируемого "
            "работника»)", "КОНТРОЛЬ")
    return missing


# ----------------------------------------------------------------------
# 7. БЫСТРЫЙ РАЗБОР ПОЛЕЙ (для проверки перечня исключений)
# ----------------------------------------------------------------------
def extract_fields_quick(text: str, sidecar: dict, model: str) -> dict:
    """Извлечь Ф.И.О./должность/подразделение/руководителя (быстрый запрос)."""
    keys = ("фио", "должность", "подразделение", "руководитель", "наставник",
            "менеджер_программ_обучения", "срок_адаптации")
    if sidecar:
        got = {k: sidecar.get(k) for k in keys if sidecar.get(k)}
        if got.get("должность"):
            log(f"поля взяты из JSON-профиля: {got}", "ПОЛЯ")
            return got
    prompt = ("Извлеки из текста профиля должности только перечисленные поля и "
              "верни СТРОГО JSON без пояснений: "
              '{"фио": "", "должность": "", "подразделение": "", '
              '"руководитель": "", "наставник": "", '
              '"менеджер_программ_обучения": ""}. '
              "Если поле не читается (рукописный текст, скан) — оставь пустую "
              "строку.\n\nТЕКСТ ПРОФИЛЯ:\n" + text[:8000])
    data = llm_json([{"role": "user", "content": prompt}], model=model, tag="ПОЛЯ")
    return {k: (data.get(k) or "") for k in keys}


# ----------------------------------------------------------------------
# 8. ГЛАВНЫЙ ПРОЦЕСС
# ----------------------------------------------------------------------
def process_card(card_path: str, model: str = MODEL, use_search: bool = True,
                 ocr_only: bool = False, force_ocr: bool = False,
                 exclude_list=None) -> dict:
    """Полный цикл по одной карточке (любой поддерживаемый формат)."""
    name = os.path.basename(card_path)
    log("=" * 70)
    log(f"КАРТОЧКА: {name} [{EXT_NAMES.get(card_ext(name), '?')}]", "ФАЙЛ")
    log("=" * 70)
    text, source = extract_text(card_path, force_ocr=force_ocr)
    log(f"источник текста: {source}; символов {len(text)}", "ФАЙЛ")
    if ocr_only:
        print("-" * 70)
        print(text[:6000])
        print("-" * 70)
        return {"файл": name, "источник_текста": source, "текст": text}
    if not text.strip():
        log("ОШИБКА: текст не извлечён (проверьте формат и качество скана)",
            "ФАЙЛ")
        return {"файл": name, "ошибка": "текст не извлечён"}
    if exclude_list is None:
        exclude_list = read_exclusion_list()
    sidecar = load_sidecar_profile(card_path)
    # --- 1) быстрый разбор полей и проверка перечня исключений ---
    fields = extract_fields_quick(text, sidecar, model)
    log(f"поля профиля: ФИО='{fields.get('фио')}', должность="
        f"'{fields.get('должность')}', подразделение="
        f"'{fields.get('подразделение')}', руководитель="
        f"'{fields.get('руководитель')}'", "ФАЙЛ")
    hit = is_excluded(fields.get("должность", ""), exclude_list)
    if hit:
        log(f"ДОЛЖНОСТЬ В ПЕРЕЧНЕ ИСКЛЮЧЕНИЙ: «{hit}» — план адаптации "
            f"НЕ формируется", "РЕШЕНИЕ")
        result = {"файл": name, "не требуется": True,
                  "должность": fields.get("должность", ""),
                  "совпадение_с_перечнем": hit, "фио": fields.get("фио", ""),
                  "подразделение": fields.get("подразделение", ""),
                  "причина": "должность включена в Приложение А РИ «Перечень "
                             "работников, которым не проводится адаптация»"}
        save_json(result, name)
        return result
    # --- 2) при нехватке данных о знаниях/умениях — профстандарт/ЕТКС ---
    position = (fields.get("должность") or "").strip()
    sparse = (len(text) < 2500
              or ("4.1" not in text and "еоретическ" not in text))
    if not position:
        log("ВНИМАНИЕ: должность не распознана (рукописная карточка?). "
            "Профстандарт по пустой должности не ищется. Чтобы поля и пробелы "
            "были точными, положите рядом JSON-профиль: "
            "«План адаптации\\_profiles\\<имя карточки>.json»", "РЕШЕНИЕ")
    elif sparse:
        log("в профиле мало данных о знаниях/умениях — ищу профстандарт/ЕТКС",
            "РЕШЕНИЕ")
    prof = profstandart_context(position,
                                enabled=use_search and sparse and bool(position))
    # --- 3) анализ и формирование плана ---
    start, end = adaptation_period(fields, text)
    log(f"срок адаптации: с {start.strftime('%d.%m.%Y')} "
        f"по {end.strftime('%d.%m.%Y')} ({(end - start).days} дней)", "СРОКИ")
    data = analyze_profile(text, sidecar, prof, start, end, model=model)
    if not data:
        log("ОШИБКА: LLM не вернула план", "ФАЙЛ")
        return {"файл": name, "ошибка": "LLM не вернула план"}
    data = apply_sidecar(data, sidecar)
    n_fix = normalize_participants(data)
    log(f"участники приведены к виду «Должность — Фамилия И.О.»: "
        f"скорректировано мероприятий {n_fix}", "АНАЛИЗ")
    data["_не_заполнено"] = audit_plan(data)
    data["файл"] = name
    data["источник_текста"] = source
    data["_профстандарт_источники"] = [s.get("href") for s in
                                       prof.get("sources", [])[:3]]
    hit2 = is_excluded(data.get("должность", ""), exclude_list)
    if hit2:
        log(f"ДОЛЖНОСТЬ В ПЕРЕЧНЕ (по ответу LLM): «{hit2}» — план не "
            f"формируется", "РЕШЕНИЕ")
        result = {"файл": name, "не требуется": True,
                  "должность": data.get("должность", ""),
                  "совпадение_с_перечнем": hit2, "фио": data.get("фио", ""),
                  "причина": "должность включена в Приложение А РИ"}
        save_json(result, name)
        return result
    log(f"пробелов выявлено: {len(data.get('пробелы', []))}; общий блок "
        f"{len(data.get('общий_блок', []))} мероприятий; индивидуальный "
        f"{len(data.get('индивидуальный_блок', []))}", "АНАЛИЗ")
    for gap in data.get("пробелы", [])[:10]:
        if isinstance(gap, dict):
            log(f"  • пробел: {gap.get('позиция')} (оценка {gap.get('оценка')} "
                f"< норматив {gap.get('норматив')})", "АНАЛИЗ")
    # --- 4) Word ---
    fio = data.get("фио") or os.path.splitext(name)[0]
    out_path = os.path.join(OUT_FOLDER,
                            f"План_адаптации_{safe_filename(fio)}.docx")
    ok = build_plan_docx(data, out_path)
    data["план_файл"] = out_path if ok else None
    save_json(data, name)
    return data


def safe_filename(text: str) -> str:
    t = re.sub(r'[\\/:*?"<>|]+', "", str(text)).strip()
    return re.sub(r"\s+", "_", t) or "Сотрудник"


def save_json(data: dict, card_name: str) -> None:
    """Сохранить результат по карточке в _cache/<карточка>_plan.json."""
    os.makedirs(CACHE_FOLDER, exist_ok=True)
    path = os.path.join(CACHE_FOLDER,
                        os.path.splitext(card_name)[0] + "_plan.json")
    clean = {k: v for k, v in data.items() if not k.startswith("_окна_dt")}
    with open(path, "w", encoding="utf-8") as f:
        json.dump(clean, f, ensure_ascii=False, indent=2, default=str)
    log(f"JSON результата: {path}", "СОХРАНЕНИЕ")


def adaptation_period(fields: dict, text: str) -> tuple:
    """Срок адаптации: из профиля, иначе — 3 месяца от сегодняшнего дня."""
    val = str((fields or {}).get("срок_адаптации") or "")
    m = re.findall(r"(\d{2})[.\-/](\d{2})[.\-/](\d{4})", val)
    if len(m) < 2:
        m = re.findall(r"(\d{2})[.\-/](\d{2})[.\-/](\d{4})", text or "")
    if len(m) >= 2:
        try:
            dates = sorted(datetime(int(y), int(mo), int(d))
                           for d, mo, y in m)
            if (dates[-1] - dates[0]).days > 30:
                return dates[0], dates[-1]
        except Exception:
            pass
    today = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
    return today, today + timedelta(days=90)


def main() -> int:
    global OLLAMA_HOST
    ap = argparse.ArgumentParser(
        description="Агент формирования плана адаптации по профилю должности")
    ap.add_argument("path", nargs="?",
                    help="файл карточки (PDF, DOCX, DOC, RTF, ODT, XLSX, XLS, "
                         "TXT, MD, PNG, JPG, TIF) или папка с карточками "
                         "(по умолчанию — папка «Сотрудники»)")
    ap.add_argument("-r", "--recursive", action="store_true",
                    help="искать карточки также во вложенных папках")
    ap.add_argument("--model", default=MODEL,
                    help=f"модель Ollama (по умолчанию {MODEL})")
    ap.add_argument("--no-search", action="store_true",
                    help="не искать профстандарт/ЕТКС в интернете")
    ap.add_argument("--ocr-only", action="store_true",
                    help="только извлечь текст карточки и показать его")
    ap.add_argument("--force-ocr", action="store_true",
                    help="игнорировать текстовый слой PDF и делать OCR (PDF)")
    ap.add_argument("--no-pause", action="store_true",
                    help="не ждать нажатия Enter в конце (автоматический запуск)")
    ap.add_argument("--host", default=OLLAMA_HOST,
                    help="адрес Ollama (например, http://192.168.1.50:11434 "
                         "для сервера в локальной сети)")
    args = ap.parse_args()

    if args.host:
        OLLAMA_HOST = args.host
    os.makedirs(OUT_FOLDER, exist_ok=True)
    log("АГЕНТ АДАПТАЦИИ: старт", "СТАРТ")
    log(f"модель: {args.model}; Ollama: "
        f"{OLLAMA_HOST or 'по умолчанию (localhost:11434)'}; поиск "
        f"профстандарта: {'выкл' if args.no_search else 'вкл'}; "
        f"thinking: {'вкл' if LLM_THINK else 'выкл'}; "
        f"OCR: EasyOCR {OCR_LANGS}; форматы: "
        f"{', '.join(dict.fromkeys(EXT_NAMES[e] for e in CARD_EXTENSIONS))}",
        "СТАРТ")
    # --- ОЧЕРЕДЬ ДОКУМЕНТОВ (любые поддерживаемые форматы) ---
    if args.path and os.path.isdir(args.path):
        queue = scan_cards(args.path, recursive=args.recursive)
    elif args.path:
        if not os.path.exists(args.path):
            log(f"ОШИБКА: файл не найден: {args.path}", "СТАРТ")
            queue = []
        elif not is_supported_card(args.path):
            log(f"ОШИБКА: формат «{card_ext(args.path) or args.path}» не "
                f"поддерживается. Поддерживаются: "
                + ", ".join(dict.fromkeys(EXT_NAMES[e] for e in CARD_EXTENSIONS)),
                "СТАРТ")
            queue = []
        else:
            log(f"в очереди 1 документ: {os.path.basename(args.path)} "
                f"[{EXT_NAMES.get(card_ext(args.path), '?')}]", "ПОИСК")
            queue = [args.path]
    else:
        queue = scan_cards(recursive=args.recursive)
    if not queue:
        log("документов для обработки нет — работа завершена", "СТАРТ")
        print("\nОБРАБОТАНО ДОКУМЕНТОВ: 0 (документы не найдены или формат "
              "не поддерживается)")
        save_log()
        wait_enter(not args.no_pause)
        return 1
    log("=" * 70, "ОЧЕРЕДЬ")
    log(f"ДОКУМЕНТОВ В ОЧЕРЕДИ: {len(queue)} — обработка ПООЧЕРЁДНО, "
        f"по одному документу за раз", "ОЧЕРЕДЬ")
    for i, card in enumerate(queue, 1):
        log(f"  {i}. {os.path.basename(card)}", "ОЧЕРЕДЬ")
    log("=" * 70, "ОЧЕРЕДЬ")
    exclude_list = read_exclusion_list()
    results, t_all = [], time.time()
    for i, card in enumerate(queue, 1):
        log(f">>> ДОКУМЕНТ {i} ИЗ {len(queue)}: {os.path.basename(card)}", "ОЧЕРЕДЬ")
        t0 = time.time()
        try:
            res = process_card(card, model=args.model,
                               use_search=not args.no_search,
                               ocr_only=args.ocr_only,
                               force_ocr=args.force_ocr,
                               exclude_list=exclude_list)
        except Exception as e:
            import traceback
            log(f"СБОЙ по карточке {os.path.basename(card)}: "
                f"{type(e).__name__}: {e}", "ОШИБКА")
            traceback.print_exc()
            res = {"файл": os.path.basename(card),
                   "ошибка": f"{type(e).__name__}: {e}"}
        res["_секунд"] = round(time.time() - t0, 1)
        results.append(res)
        log(f"<<< ДОКУМЕНТ {i} ИЗ {len(queue)} ЗАВЕРШЁН за {res['_секунд']} с; "
            f"в очереди осталось: {len(queue) - i}", "ОЧЕРЕДЬ")
    # --- СВОДКА ---
    plans = [r for r in results if r.get("план_файл")]
    skipped = [r for r in results if r.get("не требуется")]
    errors = [r for r in results if r.get("ошибка")]
    log("=" * 70, "ИТОГ")
    log(f"ОБРАБОТАНО ДОКУМЕНТОВ: {len(results)} из {len(queue)} "
        f"(общее время {round(time.time() - t_all, 1)} с)", "ИТОГ")
    log(f"  планов сформировано: {len(plans)}", "ИТОГ")
    log(f"  «не требуется» (перечень исключений): {len(skipped)}", "ИТОГ")
    log(f"  с ошибками: {len(errors)}", "ИТОГ")
    for r in results:
        if r.get("не требуется"):
            log(f"  {r.get('фио') or r.get('файл')}: НЕ ТРЕБУЕТСЯ (должность "
                f"«{r.get('должность')}» в перечне Приложения А), "
                f"{r.get('_секунд')} с", "ИТОГ")
        elif r.get("план_файл"):
            log(f"  {r.get('фио') or r.get('файл')}: план — "
                f"{os.path.basename(r['план_файл'])}, {r.get('_секунд')} с", "ИТОГ")
        else:
            log(f"  {r.get('файл')}: {r.get('ошибка') or 'без результата'}, "
                f"{r.get('_секунд')} с", "ИТОГ")
    if results:
        print("\nИТОГОВЫЙ JSON:")
        print(json.dumps(
            [{k: v for k, v in r.items()
              if k in ("файл", "фио", "должность", "не требуется",
                       "совпадение_с_перечнем", "план_файл", "пробелы")}
             for r in results], ensure_ascii=False, indent=2, default=str))
    save_log()
    print("\n" + "=" * 70)
    print(f"ОБРАБОТАНО ДОКУМЕНТОВ: {len(results)} из {len(queue)}")
    print(f"  планов сформировано:                {len(plans)}")
    print(f"  не требуется (перечень исключений): {len(skipped)}")
    print(f"  с ошибками:                         {len(errors)}")
    print(f"  папка с планами: {OUT_FOLDER}")
    print(f"  лог:             {LOG_PATH}")
    print("=" * 70)
    wait_enter(not args.no_pause)
    return 0


if __name__ == "__main__":
    sys.exit(main())








